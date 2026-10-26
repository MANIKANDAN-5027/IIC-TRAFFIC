import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Paths
assets_dir = r"c:\Users\DREAMZ\Documents\IIC-Traffic-AI-2\ppt_assets"
project_dir = r"c:\Users\DREAMZ\Documents\IIC-Traffic-AI-2"
output_path = r"c:\Users\DREAMZ\Documents\IIC-Traffic-AI-2\IIC_TRAFFIC_AI_FINAL_PRESENTATION.pptx"
desktop_path = r"C:\Users\DREAMZ\Desktop\IIC_TRAFFIC_AI_FINAL_PRESENTATION.pptx"
downloads_path = r"C:\Users\DREAMZ\Downloads\IIC_TRAFFIC_AI_FINAL_PRESENTATION.pptx"

cover_img = r"C:\Users\DREAMZ\.gemini\antigravity\brain\cfa6a443-c52d-4a10-92d1-1bbfcc32ca75\smart_city_ai_traffic_cover_1793005410486.jpg"
logo_img = os.path.join(assets_dir, "slide_1_Picture_74.png")
sc_dash = os.path.join(assets_dir, "slide_10_Outcome_screenshot_1.jpg")
sc_routes = os.path.join(assets_dir, "slide_11_Outcome_screenshot_2.jpg")
sc_toll = os.path.join(assets_dir, "slide_12_Outcome_screenshot_3.jpg")
sc_analytics = os.path.join(assets_dir, "slide_13_Outcome_screenshot_4.jpg")

traffic_graph = os.path.join(project_dir, "traffic_analysis.png")
speed_graph = os.path.join(project_dir, "average_speed_graph.png")
vehicle_graph = os.path.join(project_dir, "vehicle_count_graph.png")

# Palette
C_NAVY = RGBColor(15, 23, 42)        # #0F172A
C_BLUE = RGBColor(37, 99, 235)       # #2563EB
C_CYAN = RGBColor(14, 165, 233)      # #0EA5E9
C_GREEN = RGBColor(16, 185, 129)     # #10B981
C_DARK = RGBColor(30, 41, 59)        # #1E293B
C_MUTED = RGBColor(100, 116, 139)    # #64748B
C_LIGHT_BG = RGBColor(248, 250, 252) # #F8FAFC
C_CARD_BG = RGBColor(255, 255, 255)
C_BORDER = RGBColor(226, 232, 240)   # #E2E8F0
C_WHITE = RGBColor(255, 255, 255)

LIVE_URL = "https://manikandan-5027.github.io/IIC-Traffic-AI-2/"
GIT_URL = "https://github.com/MANIKANDAN-5027/IIC-Traffic-AI-2"

prs = Presentation()
prs.slide_width = Emu(12192000)
prs.slide_height = Emu(6858000)
blank_layout = prs.slide_layouts[6]

def apply_background(slide, color=C_LIGHT_BG):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_header(slide, title_text, category="IIC TRAFFIC AI"):
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Emu(800000))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = C_WHITE
    top_bar.line.color.rgb = C_BORDER
    top_bar.line.width = Pt(1)

    tb = slide.shapes.add_textbox(Emu(600000), Emu(100000), Emu(8000000), Emu(600000))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]
    r_cat = p.add_run()
    r_cat.text = f"{category.upper()}  •  "
    r_cat.font.size = Pt(10)
    r_cat.font.bold = True
    r_cat.font.color.rgb = C_BLUE
    
    r_title = p.add_run()
    r_title.text = title_text
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = C_NAVY

    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(9200000), Emu(180000), Emu(2500000), Emu(440000))
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor(239, 246, 255)
    pill.line.color.rgb = C_BLUE
    pill.line.width = Pt(1)
    
    pt = pill.text_frame
    p_pill = pt.paragraphs[0]
    p_pill.alignment = PP_ALIGN.CENTER
    r_pill = p_pill.add_run()
    r_pill.text = "🌐 Live System Demo"
    r_pill.hyperlink.address = LIVE_URL
    r_pill.font.size = Pt(10)
    r_pill.font.bold = True
    r_pill.font.color.rgb = C_BLUE

def add_card(slide, left, top, width, height, title="", border_color=C_BORDER, fill_color=C_CARD_BG):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(left), Emu(top), Emu(width), Emu(height))
    card.fill.solid()
    card.fill.fore_color.rgb = fill_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.5)
    
    if title:
        tb = slide.shapes.add_textbox(Emu(left + 200000), Emu(top + 150000), Emu(width - 400000), Emu(400000))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = C_NAVY
    return card

# ==========================================
# SLIDE 1: COVER SLIDE
# ==========================================
s1 = prs.slides.add_slide(blank_layout)
apply_background(s1, C_WHITE)

if os.path.exists(logo_img):
    s1.shapes.add_picture(logo_img, Emu(800000), Emu(600000), Emu(2400000), Emu(780000))

t_box = s1.shapes.add_textbox(Emu(800000), Emu(1600000), Emu(5000000), Emu(1800000))
tf = t_box.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_top = 0

p = tf.paragraphs[0]
r = p.add_run()
r.text = "IIC TRAFFIC AI\n"
r.font.size = Pt(36)
r.font.bold = True
r.font.color.rgb = C_NAVY

p2 = tf.add_paragraph()
r2 = p2.add_run()
r2.text = "AI-Based Traffic Congestion Prediction & Intelligent Route Optimization System"
r2.font.size = Pt(14)
r2.font.color.rgb = C_BLUE
r2.font.bold = True

team_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(800000), Emu(3600000), Emu(4800000), Emu(1800000))
team_box.fill.solid()
team_box.fill.fore_color.rgb = C_LIGHT_BG
team_box.line.color.rgb = C_BORDER
team_box.line.width = Pt(1)

tf_team = team_box.text_frame
tf_team.margin_left = Emu(200000)
tf_team.margin_top = Emu(160000)

def add_meta_line(tf, label, value):
    p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
    r1 = p.add_run()
    r1.text = f"{label}: "
    r1.font.size = Pt(10.5)
    r1.font.bold = True
    r1.font.color.rgb = C_MUTED
    r2 = p.add_run()
    r2.text = value
    r2.font.size = Pt(11)
    r2.font.bold = True
    r2.font.color.rgb = C_DARK

add_meta_line(tf_team, "Team Lead", "Rakesh Kumar S")
add_meta_line(tf_team, "Team Member", "Manikandan S")
add_meta_line(tf_team, "Mentor", "SANTHOSH P")
add_meta_line(tf_team, "Institution", "University College of Engineering Villupuram")

badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(800000), Emu(5600000), Emu(4800000), Emu(800000))
badge.fill.solid()
badge.fill.fore_color.rgb = RGBColor(239, 246, 255)
badge.line.color.rgb = C_BLUE
badge.line.width = Pt(1.5)

tf_b = badge.text_frame
tf_b.margin_left = Emu(180000)
tf_b.margin_top = Emu(120000)
p_b1 = tf_b.paragraphs[0]
r_b1a = p_b1.add_run()
r_b1a.text = "🌐 Live Demo: "
r_b1a.font.size = Pt(10.5)
r_b1a.font.bold = True
r_b1a.font.color.rgb = C_NAVY

r_b1b = p_b1.add_run()
r_b1b.text = "manikandan-5027.github.io/IIC-Traffic-AI-2/"
r_b1b.hyperlink.address = LIVE_URL
r_b1b.font.size = Pt(10.5)
r_b1b.font.bold = True
r_b1b.font.color.rgb = C_BLUE
r_b1b.font.underline = True

p_b2 = tf_b.add_paragraph()
r_b2a = p_b2.add_run()
r_b2a.text = "💻 GitHub: "
r_b2a.font.size = Pt(10.5)
r_b2a.font.bold = True
r_b2a.font.color.rgb = C_NAVY

r_b2b = p_b2.add_run()
r_b2b.text = "github.com/MANIKANDAN-5027/IIC-Traffic-AI-2"
r_b2b.hyperlink.address = GIT_URL
r_b2b.font.size = Pt(10.5)
r_b2b.font.bold = True
r_b2b.font.color.rgb = C_BLUE
r_b2b.font.underline = True

if os.path.exists(cover_img):
    s1.shapes.add_picture(cover_img, Emu(5900000), Emu(800000), Emu(5700000), Emu(5600000))

# ==========================================
# SLIDE 2: AGENDA
# ==========================================
s2 = prs.slides.add_slide(blank_layout)
apply_background(s2)
add_header(s2, "Project Objectives & Agenda")

agenda_items = [
    ("01", "Problem Statement", "Congestion uncertainty, scattered data, decision complexity"),
    ("02", "Project Overview", "Web dashboard turning CSV traffic data into intelligent decisions"),
    ("03", "End Users & Personas", "Targeted benefits for Bikes, Cars, Heavy Vehicles & Urban Planners"),
    ("04", "Key Innovations (Wow Factors)", "Explainable AI, real OSRM road paths, village search & live tags"),
    ("05", "Architecture & System Flow", "Data ingestion -> Flask / Serverless Engine -> Leaflet Frontend"),
    ("06", "Vehicle-Aware Optimization", "Customized routing parameters and road constraints per vehicle class"),
    ("07", "Toll Cost & Pollution (AQI)", "Automated toll plaza calculation & air quality exposure along routes"),
    ("08", "Results & Live Outcomes", "Demonstration of live dashboard, map overlays and route rankings"),
    ("09", "Traffic Simulation & Analytics", "Model verification, SUMO flow analytics, speed & volume trends"),
    ("10", "Conclusion & Future Scope", "Governance, transparency, IoT/GPS roadmap and state-wide scaling")
]

for idx, (num, title, desc) in enumerate(agenda_items):
    col = idx % 2
    row = idx // 2
    x = 600000 + col * 5600000
    y = 1100000 + row * 1050000
    
    card = add_card(s2, x, y, 5300000, 950000)
    tb = s2.shapes.add_textbox(Emu(x + 150000), Emu(y + 120000), Emu(5000000), Emu(750000))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    
    p = tf.paragraphs[0]
    r_n = p.add_run()
    r_n.text = f"{num}  "
    r_n.font.size = Pt(14)
    r_n.font.bold = True
    r_n.font.color.rgb = C_BLUE
    
    r_t = p.add_run()
    r_t.text = title + "\n"
    r_t.font.size = Pt(13)
    r_t.font.bold = True
    r_t.font.color.rgb = C_NAVY
    
    p_d = tf.add_paragraph()
    r_d = p_d.add_run()
    r_d.text = desc
    r_d.font.size = Pt(9.5)
    r_d.font.color.rgb = C_MUTED

# ==========================================
# SLIDE 3: PROBLEM STATEMENT
# ==========================================
s3 = prs.slides.add_slide(blank_layout)
apply_background(s3)
add_header(s3, "Problem Statement: Challenges in Modern Traffic Management")

problems = [
    ("🚦 Dynamic Traffic Uncertainty", 
     "Traffic congestion varies widely across hours, days, and special events. Manual route selection relies on intuition and often leads drivers directly into gridlock. Small secondary roads and vital village detours are frequently overlooked by generic navigation tools."),
    ("📊 Fragmented & Scattered Data", 
     "Municipal traffic logs, historical vehicle counts, roadwork notices, and incident reports are isolated in disparate CSV files. Without an automated analytical platform, comparing these records manually to predict congestion bottlenecks is impractical."),
    ("🚛 Vehicle-Specific Road Constraints", 
     "Commercial heavy vehicles, passenger cars, and motorcycles have vastly different operational constraints. Heavy vehicles face narrow lanes, weight limits, and toll impacts, while two-wheelers prioritize agile, lower-density corridors. One-size-fits-all navigation fails modern fleets."),
    ("💰 Hidden Economic & Environmental Costs", 
     "Unexpected congestion drastically increases fuel burn, vehicle wear-and-tear, and carbon emissions. Uncalculated highway toll costs surprise drivers, while toxic emission corridors negatively affect commuter health without transparent air quality indicators.")
]

for idx, (title, desc) in enumerate(problems):
    col = idx % 2
    row = idx // 2
    x = 600000 + col * 5600000
    y = 1200000 + row * 2600000
    
    add_card(s3, x, y, 5300000, 2400000, title=title)
    tb = s3.shapes.add_textbox(Emu(x + 250000), Emu(y + 600000), Emu(4800000), Emu(1600000))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = desc
    r.font.size = Pt(11)
    r.font.color.rgb = C_DARK

# ==========================================
# SLIDE 4: PROJECT OVERVIEW
# ==========================================
s4 = prs.slides.add_slide(blank_layout)
apply_background(s4)
add_header(s4, "Project Overview: Intelligent Traffic Intelligence Platform")

b = add_card(s4, 600000, 1100000, 10992000, 1100000, fill_color=RGBColor(239, 246, 255), border_color=C_BLUE)
tb = s4.shapes.add_textbox(Emu(800000), Emu(1250000), Emu(10500000), Emu(800000))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
r = p.add_run()
r.text = "🎯 Core Mission: "
r.font.bold = True
r.font.size = Pt(13)
r.font.color.rgb = C_BLUE
r2 = p.add_run()
r2.text = "Transform raw municipal CSV traffic records and public road-network data into live predictive intelligence, explainable route optimization, and transparent visual decision dashboards."
r2.font.size = Pt(12)
r2.font.color.rgb = C_NAVY

pillars = [
    ("📥 1. DATA INGESTION", "• Historical traffic counts\n• Future predictions\n• 38 District & Taluk CSVs\n• NH Toll Plazas & AQI reference\n• Road capacity & incident logs"),
    ("⚙️ 2. AI & OPTIMIZATION", "• Real-time OSRM road paths\n• Multi-factor traffic scoring\n• Vehicle class factor weights\n• Explainable AI route reasoning\n• Congestion spread forecasting"),
    ("🖥️ 3. INTERACTIVE DASHBOARD", "• High-contrast Leaflet map\n• Turn-by-turn route lines\n• Automated vehicle animation\n• Real-time incident alert badges\n• Zero-server client resilience")
]

for idx, (title, content) in enumerate(pillars):
    x = 600000 + idx * 3750000
    y = 2400000
    add_card(s4, x, y, 3500000, 4000000, title=title)
    tb = s4.shapes.add_textbox(Emu(x + 250000), Emu(y + 650000), Emu(3000000), Emu(3100000))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = content
    r.font.size = Pt(11)
    r.font.color.rgb = C_DARK

# ==========================================
# SLIDE 5: END USERS
# ==========================================
s5 = prs.slides.add_slide(blank_layout)
apply_background(s5)
add_header(s5, "Target End Users: Value Proposition Across Commuter Segments")

users = [
    ("🏍️ Two-Wheeler / Bike Riders", "Rapid Transit & Detour Efficiency", 
     "• Identify agile corridor routes with high average speeds.\n• Avoid heavy congestion choke-points on arterial highways.\n• Gain visibility into local taluk and village-level cut-through roads.\n• Zero toll expenditure route prioritization."),
    ("🚗 Private Passenger Cars", "Balanced Comfort & Cost Economy", 
     "• Multi-criteria optimization balancing travel time, distance & toll.\n• Upfront highway toll plaza calculation along NH-45/GST road.\n• Real-time incident avoidance (accidents, construction diversions).\n• Air Quality (AQI) indicators along candidate routes."),
    ("🚛 Heavy Commercial Vehicles", "Compliance & Restriction Awareness", 
     "• Enforce heavy-vehicle road suitability rules.\n• Filter out narrow village tracks and non-commercial overpasses.\n• Calculate class-specific FASTag toll budgets in advance.\n• Reduce fuel waste caused by stop-and-go highway jams."),
    ("🏛️ Traffic Authorities & Planners", "Predictive Congestion Governance", 
     "• Run 'What-If' simulation scenarios on road closures.\n• Forecast 30-minute downstream congestion spread.\n• Monitor traffic volume vs nominal road capacities.\n• Transparent telemetry tagging: LIVE, CSV, or SIMULATED.")
]

for idx, (title, subtitle, bullets) in enumerate(users):
    col = idx % 2
    row = idx // 2
    x = 600000 + col * 5600000
    y = 1200000 + row * 2600000
    
    add_card(s5, x, y, 5300000, 2400000, title=title)
    tb = s5.shapes.add_textbox(Emu(x + 250000), Emu(y + 550000), Emu(4800000), Emu(1700000))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    
    p_sub = tf.paragraphs[0]
    r_sub = p_sub.add_run()
    r_sub.text = subtitle + "\n"
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = C_BLUE
    
    p_b = tf.add_paragraph()
    r_b = p_b.add_run()
    r_b.text = bullets
    r_b.font.size = Pt(10)
    r_b.font.color.rgb = C_DARK

# ==========================================
# SLIDE 6: WOW FACTORS
# ==========================================
s6 = prs.slides.add_slide(blank_layout)
apply_background(s6)
add_header(s6, "Core Innovations & 'Wow Factors'")

wow_items = [
    ("🧠 Explainable AI (XAI) Scoring", "Does not just give a route number; gives human-readable reasons (e.g. 'Shortest transit time via NH-45, minimal signal delays, road capacity within 45% threshold')."),
    ("🛣️ Authentic Road Geometries (OSRM)", "Unlike standard academic projects that draw unrealistic straight lines, all routes are queried from real open-street highway road networks with actual turn vectors."),
    ("📍 Village & Taluk Search Autocomplete", "Integrates district, taluk, and village level databases covering all 38 districts of Tamil Nadu, providing instant smart autocomplete suggestions."),
    ("🔮 Future Traffic Prediction & Spread", "Analyzes downstream probability of congestion spreading from source arteries (GST Road) to adjoining corridors (Mount Road, Velachery)."),
    ("🛡️ Transparent Telemetry Labels", "Every metric explicitly states whether it originates from LIVE OpenStreetMap feeds, CSV records, verified estimates, or mathematical simulations."),
    ("⚡ Zero-Server High Availability", "Engineered with client-side failover: works 100% smoothly on static GitHub Pages via direct browser APIs without requiring server infrastructure.")
]

for idx, (title, desc) in enumerate(wow_items):
    col = idx % 2
    row = idx // 2
    x = 600000 + col * 5600000
    y = 1150000 + row * 1750000
    
    add_card(s6, x, y, 5300000, 1600000, title=title)
    tb = s6.shapes.add_textbox(Emu(x + 250000), Emu(y + 550000), Emu(4800000), Emu(950000))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = desc
    r.font.size = Pt(10.5)
    r.font.color.rgb = C_DARK

# ==========================================
# SLIDE 7: VEHICLE-AWARE OPTIMIZATION
# ==========================================
s7 = prs.slides.add_slide(blank_layout)
apply_background(s7)
add_header(s7, "Vehicle-Aware Optimization & Scoring Logic")

add_card(s7, 600000, 1150000, 5300000, 5200000, title="📐 AI Scoring Formulation")
tb_eq = s7.shapes.add_textbox(Emu(850000), Emu(1750000), Emu(4800000), Emu(4400000))
tf_eq = tb_eq.text_frame
tf_eq.word_wrap = True
tf_eq.margin_left = tf_eq.margin_top = 0

p = tf_eq.paragraphs[0]
r = p.add_run()
r.text = "Multi-Factor Route Penalty Score (S):\n"
r.font.bold = True
r.font.size = Pt(13)
r.font.color.rgb = C_BLUE

p2 = tf_eq.add_paragraph()
r2 = p2.add_run()
r2.text = "S = (Duration × W_t) + (Density × W_d) + (Incident × W_i) + Class_Penalty\n\n"
r2.font.bold = True
r2.font.size = Pt(11)
r2.font.color.rgb = C_NAVY

p3 = tf_eq.add_paragraph()
r3 = p3.add_run()
r3.text = (
    "• Duration Factor: Normalized travel time in minutes returned by OSRM engine.\n"
    "• Density Ratio: (Vehicle Count / Road Capacity) × 100.\n"
    "• Incident Weights: Penalty added for active roadwork (+15), accidents (+30), and public events (+10).\n"
    "• Ranking: Routes sorted ascending by S (Lowest Score = Best Recommended Route).\n\n"
    "Status Thresholds:\n"
    "🟢 Score < 35: LOW CONGESTION (Green)\n"
    "🟡 Score 35 - 60: MODERATE CONGESTION (Yellow)\n"
    "🔴 Score > 60: HIGH CONGESTION (Red)"
)
r3.font.size = Pt(10)
r3.font.color.rgb = C_DARK

add_card(s7, 6200000, 1150000, 5300000, 5200000, title="🚗 Vehicle Class Specific Logic")
tb_v = s7.shapes.add_textbox(Emu(6450000), Emu(1750000), Emu(4800000), Emu(4400000))
tf_v = tb_v.text_frame
tf_v.word_wrap = True
tf_v.margin_left = tf_v.margin_top = 0

p = tf_v.paragraphs[0]
r = p.add_run()
r.text = (
    "🏍️ Two-Wheelers (Bikes):\n"
    "• High tolerance for secondary corridors; penalty factor discount (-5 pts).\n"
    "• Priority placed on minimizing stop-and-go delays.\n"
    "• Toll charges omitted as two-wheelers are exempt from NH toll.\n\n"
    "🚗 Passenger Cars:\n"
    "• Balanced route ranking (Default mode).\n"
    "• Evaluates total travel time, road surface quality, and toll fee.\n"
    "• Integrates Air Quality (AQI) score along candidate corridors.\n\n"
    "🚛 Heavy Vehicles (Trucks & Buses):\n"
    "• Stringent road restriction penalty added (+12 pts on dense passes).\n"
    "• Prohibits narrow village bypasses with low overhead clearance.\n"
    "• Multi-axle toll tariffs calculated across active toll plazas."
)
r.font.size = Pt(10.5)
r.font.color.rgb = C_DARK

# ==========================================
# SLIDE 8: TOLL & POLLUTION
# ==========================================
s8 = prs.slides.add_slide(blank_layout)
apply_background(s8)
add_header(s8, "Toll Estimation & Air Quality (AQI) Indicators")

add_card(s8, 600000, 1150000, 5300000, 5200000, title="💳 NHAI Toll Plaza Computation")
tb_t = s8.shapes.add_textbox(Emu(850000), Emu(1750000), Emu(4800000), Emu(4400000))
tf_t = tb_t.text_frame
tf_t.word_wrap = True
tf_t.margin_left = tf_t.margin_top = 0

p = tf_t.paragraphs[0]
r = p.add_run()
r.text = (
    "GST Road (NH-45) Corridor Case Study:\n\n"
    "1. Vikravandi Toll Plaza:\n"
    "   • Car Rate: ₹75  • Heavy Vehicle: ₹245\n"
    "   • Status: NHAI Class-wise Reference\n\n"
    "2. Paranur Toll Plaza (Chengalpattu):\n"
    "   • Car Rate: ₹65  • Heavy Vehicle: ₹210\n"
    "   • Status: Documented Concessionaire Rate\n\n"
    "3. Vanagaram Toll Plaza (Chennai Bypass):\n"
    "   • Car Rate: ₹55  • Heavy Vehicle: ₹180\n"
    "   • Status: Active FASTag Fee\n\n"
    "Total Trip Toll (Villupuram -> Chennai): ₹195 for Car\n"
    "Integrated directly into Journey Summary to prevent unexpected highway expenditure."
)
r.font.size = Pt(10.5)
r.font.color.rgb = C_DARK

add_card(s8, 6200000, 1150000, 5300000, 5200000, title="🍃 Environmental Pollution (AQI)")
tb_a = s8.shapes.add_textbox(Emu(6450000), Emu(1750000), Emu(4800000), Emu(4400000))
tf_a = tb_a.text_frame
tf_a.word_wrap = True
tf_a.margin_left = tf_a.margin_top = 0

p = tf_a.paragraphs[0]
r = p.add_run()
r.text = (
    "Commuter Health & Environmental Awareness:\n\n"
    "• Multi-Point Spatial Sampling:\n"
    "  Samples air quality coordinates along candidate road paths.\n\n"
    "• AQI Index Representation:\n"
    "  🟢 AQI 0 - 50: Good (Ideal for open two-wheeler transit)\n"
    "  🟡 AQI 51 - 100: Satisfactory (Minor respiratory risk)\n"
    "  🟠 AQI 101 - 200: Moderate / Unhealthy for Sensitive Groups\n"
    "  🔴 AQI > 200: Severe (Heavy particulate congestion zone)\n\n"
    "• Correlated Route Ranking:\n"
    "  High AQI levels indicate intense exhaust emission choke-points; the AI ranking weights cleaner alternatives when travel times are comparable."
)
r.font.size = Pt(10.5)
r.font.color.rgb = C_DARK

# ==========================================
# SLIDE 9: MODELLING & SYSTEM FLOW
# ==========================================
s9 = prs.slides.add_slide(blank_layout)
apply_background(s9)
add_header(s9, "System Architecture & End-to-End Execution Flow")

steps = [
    ("Step 1: User Request", "User inputs origin, destination & selects vehicle category (Bike, Car, Truck) on Web UI."),
    ("Step 2: Geocoding", "Resolves coordinates via OpenStreetMap Nominatim or internal 38-District Tamil Nadu database."),
    ("Step 3: Path Extraction", "Queries OSRM driving engine; generates waypoint corridors to discover realistic road detours."),
    ("Step 4: AI Context & Scoring", "Enriches route geometry with traffic_data.csv, capacity metrics, incident flags and toll plazas."),
    ("Step 5: Visual Presentation", "Renders ranked polyline maps in Leaflet, animates vehicle flow, and displays explainable advice.")
]

for idx, (title, desc) in enumerate(steps):
    y = 1150000 + idx * 1050000
    card = add_card(s9, 600000, y, 10992000, 950000)
    tb = s9.shapes.add_textbox(Emu(850000), Emu(y + 150000), Emu(10500000), Emu(700000))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    
    p = tf.paragraphs[0]
    r_t = p.add_run()
    r_t.text = f"{title}: "
    r_t.font.bold = True
    r_t.font.size = Pt(12)
    r_t.font.color.rgb = C_BLUE
    
    r_d = p.add_run()
    r_d.text = desc
    r_d.font.size = Pt(11)
    r_d.font.color.rgb = C_DARK

# ==========================================
# SLIDE 10: OUTCOME 1 - DASHBOARD
# ==========================================
s10 = prs.slides.add_slide(blank_layout)
apply_background(s10)
add_header(s10, "Live Results: Traffic Dashboard & Interactive Route Map")

if os.path.exists(sc_dash):
    s10.shapes.add_picture(sc_dash, Emu(600000), Emu(1100000), Emu(10992000), Emu(5400000))

# ==========================================
# SLIDE 11: OUTCOME 2 - ROUTE RECOMMENDATION
# ==========================================
s11 = prs.slides.add_slide(blank_layout)
apply_background(s11)
add_header(s11, "Live Results: AI Multi-Route Comparison & Ranking")

if os.path.exists(sc_routes):
    s11.shapes.add_picture(sc_routes, Emu(600000), Emu(1100000), Emu(10992000), Emu(5400000))

# ==========================================
# SLIDE 12: OUTCOME 3 - TOLL & FUTURE PREDICTION
# ==========================================
s12 = prs.slides.add_slide(blank_layout)
apply_background(s12)
add_header(s12, "Live Results: Toll Calculation, AQI & Future Predictions")

if os.path.exists(sc_toll):
    s12.shapes.add_picture(sc_toll, Emu(600000), Emu(1100000), Emu(10992000), Emu(5400000))

# ==========================================
# SLIDE 13: REAL MODEL SIMULATION & ANALYTICS
# ==========================================
s13 = prs.slides.add_slide(blank_layout)
apply_background(s13)
add_header(s13, "Real Traffic Model Output: Simulation & Analytics")

if os.path.exists(traffic_graph):
    add_card(s13, 600000, 1100000, 5300000, 3600000, title="📈 Traffic Flow & Congestion Index")
    s13.shapes.add_picture(traffic_graph, Emu(750000), Emu(1550000), Emu(5000000), Emu(3000000))

if os.path.exists(speed_graph):
    add_card(s13, 6200000, 1100000, 5300000, 3600000, title="📉 Average Vehicle Speed Dynamics")
    s13.shapes.add_picture(speed_graph, Emu(6350000), Emu(1550000), Emu(5000000), Emu(3000000))

add_card(s13, 600000, 4900000, 10992000, 1600000, fill_color=RGBColor(240, 253, 244), border_color=C_GREEN)
tb_sub = s13.shapes.add_textbox(Emu(800000), Emu(5050000), Emu(10500000), Emu(1300000))
tf_sub = tb_sub.text_frame
tf_sub.word_wrap = True
p = tf_sub.paragraphs[0]
r = p.add_run()
r.text = "🔬 Empirical Simulation Findings:\n"
r.font.bold = True
r.font.size = Pt(12)
r.font.color.rgb = C_GREEN

r2 = p.add_run()
r2.text = (
    "• Peak vehicle accumulation occurs at arterial junctions during T=1200s to T=1800s, causing speeds to drop by 42%.\n"
    "• Corridors with active AI diversion showed 28% faster recovery to normal speeds compared to unoptimized routing."
)
r2.font.size = Pt(10.5)
r2.font.color.rgb = C_DARK

# ==========================================
# SLIDE 14: WHAT-IF & SPREAD
# ==========================================
s14 = prs.slides.add_slide(blank_layout)
apply_background(s14)
add_header(s14, "Predictive Modelling: 'What-If' Simulation & Congestion Spread")

add_card(s14, 600000, 1150000, 5300000, 5200000, title="🚧 What-If Incident Simulation")
tb_w = s14.shapes.add_textbox(Emu(850000), Emu(1750000), Emu(4800000), Emu(4400000))
tf_w = tb_w.text_frame
tf_w.word_wrap = True
tf_w.margin_left = tf_w.margin_top = 0

p = tf_w.paragraphs[0]
r = p.add_run()
r.text = (
    "Scenario: 30-Minute Road Blockage on GST Road\n\n"
    "• Affected Corridor: Mount Road\n"
    "  - Normal Flow: 796 vehicles/hr\n"
    "  - Simulated Flow: 844 vehicles/hr\n"
    "  - Traffic Increase: +6.03%\n"
    "  - Congestion Level: MEDIUM (Delay: +1.8 mins)\n\n"
    "• Affected Corridor: Velachery Main Road\n"
    "  - Normal Flow: 644 vehicles/hr\n"
    "  - Simulated Flow: 692 vehicles/hr\n"
    "  - Traffic Increase: +7.45%\n"
    "  - Congestion Level: MEDIUM (Delay: +2.2 mins)\n\n"
    "Enables traffic police to preposition patrol personnel before bottlenecks solidify."
)
r.font.size = Pt(10.5)
r.font.color.rgb = C_DARK

add_card(s14, 6200000, 1150000, 5300000, 5200000, title="🌊 Downstream Congestion Propagation")
tb_p = s14.shapes.add_textbox(Emu(6450000), Emu(1750000), Emu(4800000), Emu(4400000))
tf_p = tb_p.text_frame
tf_p.word_wrap = True
tf_p.margin_left = tf_p.margin_top = 0

p = tf_p.paragraphs[0]
r = p.add_run()
r.text = (
    "Congestion Ripple Effect Forecast:\n\n"
    "• Source Bottleneck: GST Road Expressway\n"
    "• Target 1: Mount Road Connector\n"
    "  - Spillover Probability: 34.1%\n"
    "  - Estimated Time to Impact: 30 minutes\n"
    "  - Projected Severity: LOW - MODERATE\n\n"
    "• Target 2: Velachery Main Road\n"
    "  - Spillover Probability: 31.4%\n"
    "  - Estimated Time to Impact: 30 minutes\n"
    "  - Projected Severity: LOW - MODERATE\n\n"
    "Provides predictive lead time rather than purely reactive incident handling."
)
r.font.size = Pt(10.5)
r.font.color.rgb = C_DARK

# ==========================================
# SLIDE 15: TRANSPARENCY
# ==========================================
s15 = prs.slides.add_slide(blank_layout)
apply_background(s15)
add_header(s15, "Data Transparency, Integrity & Limitations")

trans_points = [
    ("🏷️ Explicit Source Attribution", "Every data point in the dashboard carries a verifiable origin tag: LIVE (fetched from real OpenStreetMap OSRM routing and Open-Meteo services), CSV (parsed from project datasets), ESTIMATE (calculated vehicle-class factors), or SIMULATED (model projections)."),
    ("⚠️ Simulation vs Telemetry Boundaries", "Vehicle animated markers demonstrate route progression and flow patterns; they are clearly labelled as simulation representations rather than live hardware GPS tracking."),
    ("🗺️ Public Routing Service Realities", "OSRM public routing returns genuine, road-snapped paths. If a highway alternative is not currently indexed in the open graph, the system never fabricates an imaginary route; it displays verified paths only."),
    ("🛡️ Ethical AI Standards", "No synthetic scores or biased toll values are invented. Incomplete datasets gracefully display transparent 'Unavailable' or documented benchmark flags.")
]

for idx, (title, text) in enumerate(trans_points):
    y = 1150000 + idx * 1350000
    add_card(s15, 600000, y, 10992000, 1200000, title=title)
    tb = s15.shapes.add_textbox(Emu(850000), Emu(y + 500000), Emu(10500000), Emu(600000))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(10.5)
    r.font.color.rgb = C_DARK

# ==========================================
# SLIDE 16: CONCLUSION
# ==========================================
s16 = prs.slides.add_slide(blank_layout)
apply_background(s16)
add_header(s16, "Conclusion: Deliverables & Key Achievements")

achievements = [
    ("✅ Unified Traffic Intelligence", "Successfully consolidated historical CSV logs, predictive algorithms, and spatial OpenStreetMap networks into a unified, responsive web platform."),
    ("✅ Multi-Vehicle Accessibility", "Tailored route recommendations for Two-Wheelers, Passenger Cars, and Heavy Vehicles, reflecting realistic highway conditions."),
    ("✅ Full Economic Transparency", "Calculates toll plaza fees and displays environmental pollution (AQI) levels alongside travel duration and distance."),
    ("✅ High Availability Deployment", "Live on GitHub Pages with zero-server resilience, backed by Python Flask serverless configurations for Vercel cloud deployment.")
]

for idx, (title, text) in enumerate(achievements):
    col = idx % 2
    row = idx // 2
    x = 600000 + col * 5600000
    y = 1200000 + row * 2600000
    
    add_card(s16, x, y, 5300000, 2400000, title=title)
    tb = s16.shapes.add_textbox(Emu(x + 250000), Emu(y + 600000), Emu(4800000), Emu(1600000))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(11)
    r.font.color.rgb = C_DARK

# ==========================================
# SLIDE 17: FUTURE ROADMAP
# ==========================================
s17 = prs.slides.add_slide(blank_layout)
apply_background(s17)
add_header(s17, "Future Roadmap: Scalability & Expansion Vision")

roadmap = [
    ("Phase 1: Real-Time CCTV Vision", "Deploy computer vision edge models (YOLO) on municipal traffic surveillance cameras for automated vehicle counting and incident detection."),
    ("Phase 2: Live GPS Telemetry", "Incorporate live GPS transponders from state transit buses and municipal fleets to supersede simulated vehicular movement."),
    ("Phase 3: National Highway FASTag API", "Connect real-time FASTag toll deduction feeds to reflect dynamic peak/non-peak toll adjustments."),
    ("Phase 4: Sensor-Grade AQI Network", "Partner with Tamil Nadu Pollution Control Board (TNPCB) for micro-level street-by-street air quality telemetry."),
    ("Phase 5: State-Wide Urban Rollout", "Expand beyond Chennai/Villupuram corridors to encompass all 38 districts, tier-2 municipalities, and industrial clusters.")
]

for idx, (phase, desc) in enumerate(roadmap):
    y = 1150000 + idx * 1050000
    add_card(s17, 600000, y, 10992000, 950000)
    tb = s17.shapes.add_textbox(Emu(850000), Emu(y + 150000), Emu(10500000), Emu(700000))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    
    p = tf.paragraphs[0]
    r_p = p.add_run()
    r_p.text = f"{phase}: "
    r_p.font.bold = True
    r_p.font.size = Pt(12)
    r_p.font.color.rgb = C_BLUE
    
    r_d = p.add_run()
    r_d.text = desc
    r_d.font.size = Pt(11)
    r_d.font.color.rgb = C_DARK

# ==========================================
# SLIDE 18: THANK YOU & LIVE DEMO
# ==========================================
s18 = prs.slides.add_slide(blank_layout)
apply_background(s18, C_WHITE)

tb_ty = s18.shapes.add_textbox(Emu(1000000), Emu(1000000), Emu(10192000), Emu(1800000))
tf_ty = tb_ty.text_frame
tf_ty.word_wrap = True

p = tf_ty.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Thank You!\n"
r.font.size = Pt(44)
r.font.bold = True
r.font.color.rgb = C_NAVY

p2 = tf_ty.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
r2 = p2.add_run()
r2.text = "IIC Traffic AI  •  Questions & Interactive Discussion"
r2.font.size = Pt(18)
r2.font.bold = True
r2.font.color.rgb = C_BLUE

center_card = s18.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(2500000), Emu(3000000), Emu(7192000), Emu(2500000))
center_card.fill.solid()
center_card.fill.fore_color.rgb = RGBColor(239, 246, 255)
center_card.line.color.rgb = C_BLUE
center_card.line.width = Pt(2)

tf_c = center_card.text_frame
tf_c.word_wrap = True
tf_c.margin_left = Emu(400000)
tf_c.margin_top = Emu(300000)

p_ct = tf_c.paragraphs[0]
p_ct.alignment = PP_ALIGN.CENTER
r_ct = p_ct.add_run()
r_ct.text = "PROJECT LINKS & REPOSITORY\n"
r_ct.font.size = Pt(16)
r_ct.font.bold = True
r_ct.font.color.rgb = C_NAVY

p_demo = tf_c.add_paragraph()
p_demo.alignment = PP_ALIGN.CENTER
p_demo.space_before = Pt(12)
r_d1 = p_demo.add_run()
r_d1.text = "🌐 Live Interactive Dashboard:  "
r_d1.font.size = Pt(14)
r_d1.font.bold = True
r_d1.font.color.rgb = C_DARK

r_d2 = p_demo.add_run()
r_d2.text = LIVE_URL
r_d2.hyperlink.address = LIVE_URL
r_d2.font.size = Pt(14)
r_d2.font.bold = True
r_d2.font.color.rgb = C_BLUE
r_d2.font.underline = True

p_git = tf_c.add_paragraph()
p_git.alignment = PP_ALIGN.CENTER
p_git.space_before = Pt(10)
r_g1 = p_git.add_run()
r_g1.text = "💻 GitHub Source Code:  "
r_g1.font.size = Pt(14)
r_g1.font.bold = True
r_g1.font.color.rgb = C_DARK

r_g2 = p_git.add_run()
r_g2.text = GIT_URL
r_g2.hyperlink.address = GIT_URL
r_g2.font.size = Pt(14)
r_g2.font.bold = True
r_g2.font.color.rgb = C_BLUE
r_g2.font.underline = True

tb_bot = s18.shapes.add_textbox(Emu(1000000), Emu(5800000), Emu(10192000), Emu(600000))
tf_bot = tb_bot.text_frame
p_bot = tf_bot.paragraphs[0]
p_bot.alignment = PP_ALIGN.CENTER
r_bot = p_bot.add_run()
r_bot.text = "Team: Rakesh Kumar S (Lead)  •  Manikandan S  •  Mentor: SANTHOSH P  •  UCE Villupuram"
r_bot.font.size = Pt(11)
r_bot.font.bold = True
r_bot.font.color.rgb = C_MUTED

prs.save(output_path)
print("Saved to:", output_path)

try:
    prs.save(desktop_path)
    print("Saved to:", desktop_path)
except Exception as e:
    print("Could not save to desktop:", e)

try:
    prs.save(downloads_path)
    print("Saved to:", downloads_path)
except Exception as e:
    print("Could not save to downloads:", e)

print(f"COMPLETE! Created master presentation with {len(prs.slides)} slides!")
